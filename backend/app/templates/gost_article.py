"""GOST template for scientific article."""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from typing import Dict, Any

from .base import BaseTemplate


class GOSTArticleTemplate(BaseTemplate):
    """
    Template for scientific article according to GOST 7.0.7-2021.

    Key features:
    - Compact formatting
    - Smaller margins
    - Abstract and keywords sections
    - Author information
    - Article structure (Introduction, Materials and Methods, Results, Discussion, Conclusion)
    """

    # Compact formatting for articles
    FONT_SIZE = 12
    LINE_SPACING = 1.0  # Single spacing
    PARAGRAPH_INDENT = 0.75  # Smaller indent

    # Compact margins
    MARGIN_LEFT = 2.0
    MARGIN_RIGHT = 2.0
    MARGIN_TOP = 2.0
    MARGIN_BOTTOM = 2.0

    @property
    def template_name(self) -> str:
        """Template identifier."""
        return "gost_article"

    @property
    def template_description(self) -> str:
        """Human-readable template description."""
        return "ГОСТ Научная статья"

    def _create_document_styles(self, doc: Document):
        """Create document-specific styles."""
        # Article title (Heading 1) - centered, bold, larger
        self._create_heading_style(
            doc, 'Heading 1',
            font_size=14,  # Larger for article title
            bold=True,
            all_caps=False,
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
            indent=0
        )

        # Section headings (Heading 2) - bold, left-aligned
        self._create_heading_style(
            doc, 'Heading 2',
            font_size=self.FONT_SIZE,
            bold=True,
            all_caps=False,
            alignment=WD_ALIGN_PARAGRAPH.LEFT,
            indent=0
        )

        # Subsection headings (Heading 3)
        self._create_heading_style(
            doc, 'Heading 3',
            font_size=self.FONT_SIZE,
            bold=False,
            all_caps=False,
            alignment=WD_ALIGN_PARAGRAPH.LEFT,
            indent=0
        )

    def apply_formatting(self, doc: Document, structure: Dict[str, Any]) -> Document:
        """
        Apply article formatting to the document.

        Args:
            doc: Document to format
            structure: Document structure from AI analysis

        Returns:
            Formatted document
        """
        # Apply base formatting
        self._apply_page_settings(doc)
        self._create_base_styles(doc)
        self._create_document_styles(doc)

        # Add article metadata
        self._add_article_metadata(doc, structure)

        # Process document sections
        if 'sections' in structure:
            for section in structure['sections']:
                self._format_section(doc, section)

        return doc

    def _add_article_metadata(self, doc: Document, structure: Dict[str, Any]):
        """
        Add article-specific metadata.

        Includes:
        - Title
        - Authors
        - Affiliations
        - Abstract
        - Keywords
        """
        # Title (if provided in structure)
        if 'title' in structure:
            title_para = doc.add_paragraph()
            title_run = title_para.add_run(structure['title'])
            title_run.font.size = Pt(14)
            title_run.font.bold = True
            title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            title_para.paragraph_format.space_after = Pt(12)

        # Authors (if provided)
        if 'authors' in structure:
            authors_para = doc.add_paragraph()
            authors_run = authors_para.add_run(', '.join(structure['authors']))
            authors_run.font.size = Pt(self.FONT_SIZE)
            authors_run.font.italic = True
            authors_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            authors_para.paragraph_format.space_after = Pt(6)

        # Affiliations (if provided)
        if 'affiliations' in structure:
            affil_para = doc.add_paragraph()
            affil_run = affil_para.add_run(', '.join(structure['affiliations']))
            affil_run.font.size = Pt(10)
            affil_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            affil_para.paragraph_format.space_after = Pt(12)

        # Abstract section
        abstract_heading = doc.add_paragraph()
        abstract_heading_run = abstract_heading.add_run("Аннотация")
        abstract_heading_run.font.bold = True
        abstract_heading_run.font.size = Pt(self.FONT_SIZE)
        abstract_heading.paragraph_format.space_after = Pt(6)

        if 'abstract' in structure:
            abstract_para = doc.add_paragraph(structure['abstract'])
            abstract_para.paragraph_format.space_after = Pt(12)

        # Keywords section
        if 'keywords' in structure:
            keywords_para = doc.add_paragraph()
            keywords_label = keywords_para.add_run("Ключевые слова: ")
            keywords_label.font.bold = True
            keywords_label.font.size = Pt(self.FONT_SIZE)

            keywords_text = keywords_para.add_run(', '.join(structure['keywords']))
            keywords_text.font.size = Pt(self.FONT_SIZE)

            keywords_para.paragraph_format.space_after = Pt(18)

    def _format_section(self, doc: Document, section: Dict[str, Any]):
        """Format a single section."""
        # Add heading
        heading = doc.add_heading(section.get('title', ''), level=section.get('level', 2))

        # Add content paragraphs
        if 'content' in section:
            for content_item in section['content']:
                if content_item.get('type') == 'paragraph':
                    para = doc.add_paragraph(content_item.get('text', ''))

        # Process subsections recursively
        if 'subsections' in section:
            for subsection in section['subsections']:
                self._format_section(doc, subsection)
