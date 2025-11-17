from docx import Document
from typing import Dict, Any, Tuple
import os
from app.services.document_analyzer import DocumentAnalyzer
from app.templates.factory import TemplateFactory
from app.templates.base import BaseTemplate


class DocxProcessor:
    """Основной процессор для обработки DOCX документов"""

    def __init__(self, template_name: str = "gost_vkr"):
        """
        Инициализация процессора

        Args:
            template_name: Название шаблона форматирования
        """
        self.analyzer = DocumentAnalyzer()
        self.template = TemplateFactory.create_template(template_name)
        self.template_name = template_name

    def extract_text(self, doc: Document) -> str:
        """
        Извлекает текст из документа для анализа

        Args:
            doc: Объект документа python-docx

        Returns:
            Текстовое содержимое документа
        """

        full_text = []

        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                full_text.append(paragraph.text)

        return '\n'.join(full_text)

    async def process_document(
        self,
        input_path: str,
        output_path: str
    ) -> Tuple[str, Dict[str, Any]]:
        """
        Полный цикл обработки документа:
        1. Открытие документа
        2. Извлечение текста
        3. Анализ структуры с помощью AI
        4. Применение форматирования по шаблону
        5. Сохранение результата

        Args:
            input_path: Путь к исходному файлу
            output_path: Путь для сохранения результата

        Returns:
            Кортеж (путь к результату, структура документа)
        """

        try:
            # Открываем документ
            doc = Document(input_path)

            # Извлекаем текст для анализа
            text_content = self.extract_text(doc)

            if not text_content.strip():
                raise ValueError("Документ пустой или не содержит текста")

            # Анализируем структуру с помощью AI
            structure = await self.analyzer.analyze_document_structure(text_content)

            # Применяем форматирование согласно выбранному шаблону
            formatted_doc = self.template.apply_formatting(doc, structure)

            # Сохраняем результат
            formatted_doc.save(output_path)

            return output_path, structure

        except Exception as e:
            raise Exception(f"Ошибка обработки документа: {str(e)}")

    def validate_docx(self, file_path: str) -> bool:
        """
        Проверяет, является ли файл корректным DOCX документом

        Args:
            file_path: Путь к файлу

        Returns:
            True если файл корректный, иначе False
        """

        try:
            doc = Document(file_path)
            # Проверяем, что документ содержит хотя бы один параграф
            return len(doc.paragraphs) > 0
        except Exception:
            return False

    def get_document_stats(self, doc: Document) -> Dict[str, Any]:
        """
        Получает статистику документа

        Args:
            doc: Объект документа python-docx

        Returns:
            Словарь со статистикой
        """

        paragraphs_count = len(doc.paragraphs)
        tables_count = len(doc.tables)

        # Подсчет слов
        words_count = 0
        for paragraph in doc.paragraphs:
            words_count += len(paragraph.text.split())

        return {
            "paragraphs": paragraphs_count,
            "tables": tables_count,
            "words": words_count,
        }
