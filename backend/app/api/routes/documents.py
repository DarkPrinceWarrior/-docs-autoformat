from fastapi import APIRouter, UploadFile, File, HTTPException, Depends, Form
from fastapi.responses import FileResponse, JSONResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select
from app.core.database import get_db
from app.models.document import Document, DocumentStatus
from app.tasks.document_tasks import process_document_task
from app.core.config import settings
from app.templates.factory import TemplateFactory
from app.services.gost_validator import GOSTValidator
from app.services.toc_generator import TOCGenerator
from app.services.bibliography_formatter import BibliographyFormatter
from app.services.figure_table_numbering import FigureTableNumbering
from typing import List, Optional, Dict, Any
from pydantic import BaseModel
import os
import shutil
import uuid

router = APIRouter()


class DocumentResponse(BaseModel):
    """Модель ответа с информацией о документе"""
    id: int
    original_filename: str
    template_name: str
    status: str
    created_at: str
    structure_analysis: str | None = None
    error_message: str | None = None

    class Config:
        from_attributes = True


class TemplateInfo(BaseModel):
    """Информация о шаблоне форматирования"""
    name: str
    description: str
    font: str
    font_size: str
    line_spacing: str


class TaskStatusResponse(BaseModel):
    """Модель ответа со статусом задачи"""
    task_id: str
    status: str
    document_id: int | None = None


@router.get("/templates", response_model=List[TemplateInfo])
async def get_available_templates():
    """
    Получение списка доступных шаблонов форматирования

    Возвращает информацию о всех доступных шаблонах
    """
    templates = TemplateFactory.get_available_templates()
    return templates


@router.post("/upload", response_model=DocumentResponse)
async def upload_document(
    file: UploadFile = File(...),
    template: str = Form("gost_vkr"),
    db: AsyncSession = Depends(get_db)
):
    """
    Загрузка документа для обработки

    - **file**: DOCX файл для обработки
    - **template**: Название шаблона форматирования (по умолчанию: gost_vkr)
    """

    # Проверяем, что шаблон существует
    try:
        TemplateFactory.create_template(template)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))

    # Проверяем расширение файла
    if not file.filename.endswith('.docx'):
        raise HTTPException(
            status_code=400,
            detail="Поддерживаются только файлы формата .docx"
        )

    # Проверяем размер файла
    file.file.seek(0, 2)
    file_size = file.file.tell()
    file.file.seek(0)

    if file_size > settings.MAX_FILE_SIZE:
        raise HTTPException(
            status_code=400,
            detail=f"Размер файла превышает максимально допустимый ({settings.MAX_FILE_SIZE / 1024 / 1024} МБ)"
        )

    # Создаем директорию для загрузок, если её нет
    os.makedirs(settings.UPLOAD_DIR, exist_ok=True)

    # Генерируем уникальное имя файла
    file_id = str(uuid.uuid4())
    file_path = os.path.join(settings.UPLOAD_DIR, f"{file_id}_{file.filename}")

    # Сохраняем файл
    try:
        with open(file_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)
    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"Ошибка при сохранении файла: {str(e)}"
        )

    # Создаем запись в базе данных
    document = Document(
        original_filename=file.filename,
        original_file_path=file_path,
        template_name=template,
        status=DocumentStatus.UPLOADED
    )

    db.add(document)
    await db.commit()
    await db.refresh(document)

    # Запускаем задачу обработки
    task = process_document_task.delay(document.id)

    return DocumentResponse(
        id=document.id,
        original_filename=document.original_filename,
        template_name=document.template_name,
        status=document.status.value,
        created_at=str(document.created_at),
        structure_analysis=document.structure_analysis,
        error_message=document.error_message
    )


@router.get("/documents/{document_id}", response_model=DocumentResponse)
async def get_document_status(
    document_id: int,
    db: AsyncSession = Depends(get_db)
):
    """
    Получение статуса обработки документа

    - **document_id**: ID документа
    """

    result = await db.execute(
        select(Document).where(Document.id == document_id)
    )
    document = result.scalar_one_or_none()

    if not document:
        raise HTTPException(status_code=404, detail="Документ не найден")

    return DocumentResponse(
        id=document.id,
        original_filename=document.original_filename,
        template_name=document.template_name,
        status=document.status.value,
        created_at=str(document.created_at),
        structure_analysis=document.structure_analysis,
        error_message=document.error_message
    )


@router.get("/documents/{document_id}/download")
async def download_formatted_document(
    document_id: int,
    db: AsyncSession = Depends(get_db)
):
    """
    Скачивание отформатированного документа

    - **document_id**: ID документа
    """

    result = await db.execute(
        select(Document).where(Document.id == document_id)
    )
    document = result.scalar_one_or_none()

    if not document:
        raise HTTPException(status_code=404, detail="Документ не найден")

    if document.status != DocumentStatus.COMPLETED:
        raise HTTPException(
            status_code=400,
            detail=f"Документ еще не обработан. Текущий статус: {document.status.value}"
        )

    if not document.formatted_file_path or not os.path.exists(document.formatted_file_path):
        raise HTTPException(
            status_code=404,
            detail="Отформатированный файл не найден"
        )

    filename = f"formatted_{document.original_filename}"
    return FileResponse(
        path=document.formatted_file_path,
        filename=filename,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )


@router.get("/documents", response_model=List[DocumentResponse])
async def list_documents(
    skip: int = 0,
    limit: int = 100,
    db: AsyncSession = Depends(get_db)
):
    """
    Получение списка всех документов

    - **skip**: Количество пропускаемых записей
    - **limit**: Максимальное количество возвращаемых записей
    """

    result = await db.execute(
        select(Document)
        .order_by(Document.created_at.desc())
        .offset(skip)
        .limit(limit)
    )
    documents = result.scalars().all()

    return [
        DocumentResponse(
            id=doc.id,
            original_filename=doc.original_filename,
            template_name=doc.template_name,
            status=doc.status.value,
            created_at=str(doc.created_at),
            structure_analysis=doc.structure_analysis,
            error_message=doc.error_message
        )
        for doc in documents
    ]


@router.post("/documents/{document_id}/validate")
async def validate_document(
    document_id: int,
    db: AsyncSession = Depends(get_db)
) -> Dict[str, Any]:
    """
    Валидация документа на соответствие ГОСТ

    - **document_id**: ID документа для валидации
    """
    result = await db.execute(
        select(Document).where(Document.id == document_id)
    )
    document = result.scalar_one_or_none()

    if not document:
        raise HTTPException(status_code=404, detail="Документ не найден")

    if document.status != DocumentStatus.COMPLETED:
        raise HTTPException(
            status_code=400,
            detail="Документ еще не обработан. Валидация доступна только для завершенных документов"
        )

    if not document.formatted_file_path or not os.path.exists(document.formatted_file_path):
        raise HTTPException(
            status_code=404,
            detail="Отформатированный файл не найден"
        )

    # Validate document
    validator = GOSTValidator(template_name=document.template_name)
    validation_result = validator.validate_document(document.formatted_file_path)

    return validation_result


@router.post("/documents/{document_id}/generate-toc")
async def generate_toc(
    document_id: int,
    db: AsyncSession = Depends(get_db)
) -> Dict[str, Any]:
    """
    Генерация оглавления для документа

    - **document_id**: ID документа
    """
    result = await db.execute(
        select(Document).where(Document.id == document_id)
    )
    document = result.scalar_one_or_none()

    if not document:
        raise HTTPException(status_code=404, detail="Документ не найден")

    if document.status != DocumentStatus.COMPLETED:
        raise HTTPException(
            status_code=400,
            detail="Генерация TOC доступна только для завершенных документов"
        )

    if not document.formatted_file_path or not os.path.exists(document.formatted_file_path):
        raise HTTPException(
            status_code=404,
            detail="Отформатированный файл не найден"
        )

    # Generate TOC structure
    toc_generator = TOCGenerator(template_name=document.template_name)

    from docx import Document as DocxDocument
    doc = DocxDocument(document.formatted_file_path)

    toc_structure = toc_generator.get_toc_structure(doc)

    return {
        'document_id': document_id,
        'toc': toc_structure,
        'message': 'Оглавление сгенерировано. Используйте /insert-toc для вставки в документ'
    }


@router.post("/documents/{document_id}/validate-bibliography")
async def validate_bibliography(
    document_id: int,
    db: AsyncSession = Depends(get_db)
) -> Dict[str, Any]:
    """
    Проверка списка литературы

    - **document_id**: ID документа
    """
    result = await db.execute(
        select(Document).where(Document.id == document_id)
    )
    document = result.scalar_one_or_none()

    if not document:
        raise HTTPException(status_code=404, detail="Документ не найден")

    if document.status != DocumentStatus.COMPLETED:
        raise HTTPException(
            status_code=400,
            detail="Проверка библиографии доступна только для завершенных документов"
        )

    if not document.formatted_file_path or not os.path.exists(document.formatted_file_path):
        raise HTTPException(
            status_code=404,
            detail="Отформатированный файл не найден"
        )

    # Validate bibliography
    bib_formatter = BibliographyFormatter(template_name=document.template_name)

    from docx import Document as DocxDocument
    doc = DocxDocument(document.formatted_file_path)

    validation_result = bib_formatter.validate_bibliography(doc)

    return {
        'document_id': document_id,
        'validation': validation_result
    }


@router.post("/documents/{document_id}/validate-numbering")
async def validate_numbering(
    document_id: int,
    db: AsyncSession = Depends(get_db)
) -> Dict[str, Any]:
    """
    Проверка нумерации таблиц и рисунков

    - **document_id**: ID документа
    """
    result = await db.execute(
        select(Document).where(Document.id == document_id)
    )
    document = result.scalar_one_or_none()

    if not document:
        raise HTTPException(status_code=404, detail="Документ не найден")

    if document.status != DocumentStatus.COMPLETED:
        raise HTTPException(
            status_code=400,
            detail="Проверка нумерации доступна только для завершенных документов"
        )

    if not document.formatted_file_path or not os.path.exists(document.formatted_file_path):
        raise HTTPException(
            status_code=404,
            detail="Отформатированный файл не найден"
        )

    # Validate numbering
    numbering = FigureTableNumbering(template_name=document.template_name)

    from docx import Document as DocxDocument
    doc = DocxDocument(document.formatted_file_path)

    validation_result = numbering.validate_numbering(doc)

    return {
        'document_id': document_id,
        'validation': validation_result
    }
