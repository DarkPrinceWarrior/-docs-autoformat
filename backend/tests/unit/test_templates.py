"""Unit tests for template system."""
import pytest
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.templates.factory import TemplateFactory
from app.templates.gost_vkr import GOSTVKRTemplate
from app.templates.gost_coursework import GOSTCourseworkTemplate
from app.templates.base import BaseTemplate


class TestTemplateFactory:
    """Tests for TemplateFactory."""

    def test_get_available_templates(self):
        """Test that get_available_templates returns correct list."""
        templates = TemplateFactory.get_available_templates()

        assert len(templates) == 2
        assert any(t['name'] == 'gost_vkr' for t in templates)
        assert any(t['name'] == 'gost_coursework' for t in templates)

        # Check structure
        for template in templates:
            assert 'name' in template
            assert 'description' in template
            assert 'font' in template
            assert 'font_size' in template
            assert 'line_spacing' in template

    def test_create_template_vkr(self):
        """Test creating GOST VKR template."""
        template = TemplateFactory.create_template('gost_vkr')

        assert isinstance(template, GOSTVKRTemplate)
        assert template.template_name == 'gost_vkr'

    def test_create_template_coursework(self):
        """Test creating GOST Coursework template."""
        template = TemplateFactory.create_template('gost_coursework')

        assert isinstance(template, GOSTCourseworkTemplate)
        assert template.template_name == 'gost_coursework'

    def test_create_template_invalid(self):
        """Test that invalid template name raises ValueError."""
        with pytest.raises(ValueError, match="Шаблон 'invalid' не найден"):
            TemplateFactory.create_template('invalid')


class TestGOSTVKRTemplate:
    """Tests for GOST VKR template."""

    @pytest.fixture
    def template(self):
        """Create template instance."""
        return GOSTVKRTemplate()

    @pytest.fixture
    def sample_structure(self):
        """Create sample document structure."""
        return {
            "sections": [
                {
                    "title": "Введение",
                    "level": 1,
                    "content": [
                        {"type": "paragraph", "text": "Это первый абзац."},
                        {"type": "paragraph", "text": "Это второй абзац."}
                    ]
                },
                {
                    "title": "Основная часть",
                    "level": 1,
                    "content": [
                        {"type": "paragraph", "text": "Содержание раздела."}
                    ],
                    "subsections": [
                        {
                            "title": "Подраздел 1.1",
                            "level": 2,
                            "content": [
                                {"type": "paragraph", "text": "Содержание подраздела."}
                            ]
                        }
                    ]
                }
            ]
        }

    def test_template_name(self, template):
        """Test template name property."""
        assert template.template_name == 'gost_vkr'

    def test_template_description(self, template):
        """Test template description property."""
        assert 'ВКР' in template.template_description
        assert 'ГОСТ' in template.template_description

    def test_constants(self, template):
        """Test template constants."""
        assert template.FONT_NAME == "Times New Roman"
        assert template.FONT_SIZE == 14
        assert template.LINE_SPACING == 1.5
        assert template.PARAGRAPH_INDENT == 1.25
        assert template.MARGIN_LEFT == 3.0
        assert template.MARGIN_RIGHT == 1.0
        assert template.MARGIN_TOP == 2.0
        assert template.MARGIN_BOTTOM == 2.0

    def test_apply_formatting(self, template, sample_structure):
        """Test that apply_formatting creates a properly formatted document."""
        doc = Document()
        formatted_doc = template.apply_formatting(doc, sample_structure)

        assert isinstance(formatted_doc, Document)
        assert len(formatted_doc.paragraphs) > 0


class TestGOSTCourseworkTemplate:
    """Tests for GOST Coursework template."""

    @pytest.fixture
    def template(self):
        """Create template instance."""
        return GOSTCourseworkTemplate()

    def test_template_name(self, template):
        """Test template name property."""
        assert template.template_name == 'gost_coursework'

    def test_template_description(self, template):
        """Test template description property."""
        assert 'Курсовая' in template.template_description
        assert 'ГОСТ' in template.template_description

    def test_margin_difference(self, template):
        """Test that coursework has wider right margin."""
        assert template.MARGIN_RIGHT == 1.5

        vkr_template = GOSTVKRTemplate()
        assert template.MARGIN_RIGHT > vkr_template.MARGIN_RIGHT

    def test_constants_inheritance(self, template):
        """Test that template inherits base constants."""
        # Should inherit from BaseTemplate
        assert template.FONT_NAME == "Times New Roman"
        assert template.FONT_SIZE == 14
        assert template.LINE_SPACING == 1.5
        assert template.PARAGRAPH_INDENT == 1.25


class TestBaseTemplate:
    """Tests for base template methods."""

    class ConcreteTemplate(BaseTemplate):
        """Concrete implementation for testing."""

        @property
        def template_name(self):
            return "test_template"

        @property
        def template_description(self):
            return "Test Template"

        def apply_formatting(self, doc, structure):
            self._apply_page_settings(doc)
            self._create_base_styles(doc)
            return doc

    @pytest.fixture
    def template(self):
        """Create concrete template instance."""
        return self.ConcreteTemplate()

    def test_apply_page_settings(self, template):
        """Test page settings application."""
        doc = Document()
        template._apply_page_settings(doc)

        section = doc.sections[0]

        # Check margins (convert to cm for comparison)
        assert section.left_margin.cm == pytest.approx(template.MARGIN_LEFT, rel=0.1)
        assert section.right_margin.cm == pytest.approx(template.MARGIN_RIGHT, rel=0.1)
        assert section.top_margin.cm == pytest.approx(template.MARGIN_TOP, rel=0.1)
        assert section.bottom_margin.cm == pytest.approx(template.MARGIN_BOTTOM, rel=0.1)

    def test_create_base_styles(self, template):
        """Test base styles creation."""
        doc = Document()
        template._create_base_styles(doc)

        # Check normal style exists and has correct properties
        normal_style = doc.styles['Normal']
        assert normal_style.font.name == template.FONT_NAME
        assert normal_style.font.size.pt == template.FONT_SIZE
        assert normal_style.paragraph_format.line_spacing == template.LINE_SPACING
        assert normal_style.paragraph_format.first_line_indent.cm == pytest.approx(
            template.PARAGRAPH_INDENT, rel=0.1
        )
        assert normal_style.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY

    def test_create_heading_style(self, template):
        """Test heading style creation."""
        doc = Document()

        template._create_heading_style(
            doc,
            'TestHeading1',
            font_size=16,
            bold=True,
            all_caps=True,
            alignment=WD_ALIGN_PARAGRAPH.CENTER
        )

        style = doc.styles['TestHeading1']
        assert style.font.size.pt == 16
        assert style.font.bold is True
        assert style.font.all_caps is True
        assert style.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
