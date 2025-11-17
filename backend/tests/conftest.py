"""Pytest configuration and fixtures."""
import pytest
import os
import tempfile
from pathlib import Path
from typing import Generator
from httpx import AsyncClient
from sqlalchemy.ext.asyncio import create_async_engine, AsyncSession
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.main import app
from app.core.database import Base, get_db
from app.core.config import settings


# Test database URL (in-memory SQLite for tests)
TEST_DATABASE_URL = "sqlite+aiosqlite:///:memory:"


@pytest.fixture(scope="session")
def event_loop():
    """Create an event loop for the test session."""
    import asyncio
    loop = asyncio.get_event_loop_policy().new_event_loop()
    yield loop
    loop.close()


@pytest.fixture
async def test_engine():
    """Create test database engine."""
    engine = create_async_engine(
        TEST_DATABASE_URL,
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )

    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

    yield engine

    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.drop_all)

    await engine.dispose()


@pytest.fixture
async def test_db(test_engine) -> AsyncSession:
    """Create test database session."""
    async_session = sessionmaker(
        test_engine, class_=AsyncSession, expire_on_commit=False
    )

    async with async_session() as session:
        yield session


@pytest.fixture
async def client(test_db) -> Generator:
    """Create test client with database override."""

    async def override_get_db():
        yield test_db

    app.dependency_overrides[get_db] = override_get_db

    async with AsyncClient(app=app, base_url="http://test") as ac:
        yield ac

    app.dependency_overrides.clear()


@pytest.fixture
def temp_dir():
    """Create temporary directory for test files."""
    with tempfile.TemporaryDirectory() as tmpdirname:
        yield Path(tmpdirname)


@pytest.fixture
def sample_docx_path(temp_dir):
    """Create a sample DOCX file for testing."""
    from docx import Document

    doc = Document()
    doc.add_heading('Test Document', 0)
    doc.add_paragraph('This is a test paragraph.')

    file_path = temp_dir / 'test_document.docx'
    doc.save(str(file_path))

    return file_path


@pytest.fixture
def mock_ai_provider(mocker):
    """Mock AI provider for testing."""
    mock_provider = mocker.Mock()
    mock_provider.analyze_document_structure = mocker.AsyncMock(
        return_value={
            "sections": [
                {
                    "title": "Test Section",
                    "level": 1,
                    "content": [
                        {"type": "paragraph", "text": "Test content"}
                    ]
                }
            ]
        }
    )
    return mock_provider
